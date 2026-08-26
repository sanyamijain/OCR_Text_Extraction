from __future__ import annotations

import json
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "output" / "surya"
DESTINATION = ROOT / "output" / "documents"

def text_of(html: str) -> str:
    return BeautifulSoup(html or "", "html.parser").get_text(" ", strip=True)

def main() -> int:
    files = sorted(SOURCE.glob("*_surya.json"))
    if not files:
        print("No Surya JSON output found. Run run_surya.py first.")
        return 1
    DESTINATION.mkdir(parents=True, exist_ok=True)
    for source in files:
        pages = json.loads(source.read_text(encoding="utf-8"))
        stem = source.stem.removesuffix("_surya")
        doc = Document()
        doc.add_heading(stem, level=0)
        wb = Workbook()
        ws = wb.active
        ws.title = "Surya Blocks"
        ws.append(["Page", "Order", "Type", "Text", "Bounding box"])
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
        for page in pages:
            page_number = page.get("source_page_index", 0) + 1
            doc.add_heading(f"Page {page_number}", level=1)
            blocks = sorted(page.get("blocks", []), key=lambda b: b.get("reading_order", 0))
            for block in blocks:
                text = text_of(block.get("html", ""))
                if not text:
                    continue
                paragraph = doc.add_paragraph(text)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                ws.append([page_number, block.get("reading_order"), block.get("label"), text,
                           json.dumps(block.get("bbox", block.get("polygon", [])))])
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 10
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 100
        ws.column_dimensions["E"].width = 45
        for row in ws.iter_rows(min_row=2):
            row[3].alignment = Alignment(wrap_text=True, vertical="top")
        doc.save(DESTINATION / f"{stem}.docx")
        wb.save(DESTINATION / f"{stem}.xlsx")
        print(f"Created Word and Excel files for {stem}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())

