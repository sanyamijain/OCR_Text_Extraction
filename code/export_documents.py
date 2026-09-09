from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from document_parser import numbered_items, parse_questions, section_kind, text_of
from excel_math import readable_math
from word_math import WordMathRenderer
from structured_output import MATH_MARKER, html_logical_lines

CODE_ROOT = Path(__file__).resolve().parent
ROOT = CODE_ROOT.parent
SOURCE = ROOT / "output" / "surya"
DESTINATION = ROOT / "output" / "documents"
EVALUATION = ROOT / "reports" / "evaluation"
# The exporter uses Unicode script evidence rather than a fixed Hindi-only
# schema.  This allows a workbook to add Bengali, Tamil, Arabic, etc. columns
# when those scripts occur in the extracted question content.
SCRIPT_COLUMNS = (
    ("English", r"[A-Za-z]"),
    ("हिंदी", r"[\u0900-\u097F]"),
    ("বাংলা", r"[\u0980-\u09FF]"),
    ("ગુજરાતી", r"[\u0A80-\u0AFF]"),
    ("ਪੰਜਾਬੀ", r"[\u0A00-\u0A7F]"),
    ("தமிழ்", r"[\u0B80-\u0BFF]"),
    ("తెలుగు", r"[\u0C00-\u0C7F]"),
    ("ಕನ್ನಡ", r"[\u0C80-\u0CFF]"),
    ("മലയാളം", r"[\u0D00-\u0D7F]"),
    ("ଓଡ଼ିଆ", r"[\u0B00-\u0B7F]"),
    ("العربية", r"[\u0600-\u06FF]"),
)


def language_values(text: str) -> dict[str, str]:
    """Split recognized text by script for editable Excel cells; no translation."""
    values: dict[str, str] = {}
    # A script change normally marks a printed bilingual boundary. New lines
    # are also preserved as separate OCR reading-order units.
    pieces = re.split(
        r"\n+|(?<=[A-Za-z])\s+(?=[\u0600-\u06FF\u0900-\u0D7F])|"
        r"(?<=[\u0600-\u06FF\u0900-\u0D7F])\s+(?=[A-Za-z])",
        text or "",
    )
    for piece in pieces:
        piece = re.sub(r"\s+", " ", piece).strip()
        if not piece:
            continue
        language, count = max(
            ((name, len(re.findall(pattern, piece))) for name, pattern in SCRIPT_COLUMNS),
            key=lambda item: item[1],
        )
        if count:
            values[language] = f"{values.get(language, '')} {piece}".strip()
    return values


def detected_languages(items) -> list[str]:
    """Return all scripts used in one section, in stable spreadsheet order."""
    found: set[str] = set()
    for question in items:
        values = [question.english, question.hindi]
        values.extend(question.english_options.values())
        values.extend(question.hindi_options.values())
        for value in values:
            found.update(language_values(value).keys())
    return [name for name, _ in SCRIPT_COLUMNS if name in found]


def question_language_values(question) -> dict[str, str]:
    """Merge parser fields while retaining each detected script separately."""
    result: dict[str, str] = {}
    for value in (question.english, question.hindi):
        for language, content in language_values(value).items():
            result[language] = f"{result.get(language, '')} {content}".strip()
    return result


def option_language_values(question, letter: str) -> dict[str, str]:
    """Merge the parser's bilingual option fields without losing either text."""
    result: dict[str, str] = {}
    for value in (question.english_options.get(letter, ""), question.hindi_options.get(letter, "")):
        for language, content in language_values(value).items():
            result[language] = f"{result.get(language, '')} {content}".strip()
    return result


def safe_sheet_name(name: str, existing: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", name)[:31] or "Questions"
    candidate = base
    index = 2
    while candidate in existing:
        suffix = f" ({index})"
        candidate = base[: 31 - len(suffix)] + suffix
        index += 1
    existing.add(candidate)
    return candidate


def style_header(ws, row_number: int, column_count: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[row_number][:column_count]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_sheet(ws, header_row: int, widths: list[int]) -> None:
    style_header(ws, header_row, len(widths))
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{ws.max_row}"
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_number in range(header_row + 1, ws.max_row + 1):
        ws.row_dimensions[row_number].height = 42


def cover_lines(pages: list[dict]) -> list[str]:
    if not pages:
        return []
    values: list[str] = []
    for block in sorted(pages[0].get("blocks", []), key=lambda b: b.get("reading_order", 0)):
        if block.get("label") in {"PageHeader", "PageFooter", "Picture"} or block.get("skipped"):
            continue
        text = text_of(block.get("html", ""))
        lower = text.lower()
        if not text or "instructions for" in lower or "परीक्षार्थियों के लिये" in text:
            break
        if "booklet set code" in lower or "प्रश्न पुस्तिका सेट कोड" in text:
            continue
        values.append(readable_math(text))
    return values


def section_intros(pages: list[dict]) -> dict[str, list[str]]:
    intros: dict[str, list[str]] = defaultdict(list)
    current: str | None = None
    started: set[str] = set()
    pending_heading = ""
    for page in pages:
        for block in sorted(page.get("blocks", []), key=lambda b: b.get("reading_order", 0)):
            if block.get("label") in {"PageHeader", "PageFooter", "Picture"} or block.get("skipped"):
                continue
            text = text_of(block.get("html", ""))
            if not text:
                continue
            if block.get("label") == "SectionHeader":
                detected = section_kind(text)
                if detected:
                    current = detected
                    if pending_heading:
                        intros[current].append(pending_heading)
                    intros[current].append(text)
                    pending_heading = ""
                    continue
                if "section" in text.lower() or "खण्ड" in text:
                    pending_heading = text
            if current and current not in started:
                if numbered_items(text):
                    started.add(current)
                else:
                    intros[current].append(readable_math(text))
    return intros


def add_intro(ws, lines: list[str], columns: int) -> int:
    if not lines:
        return 1
    for index, line in enumerate(lines):
        row = ws.max_row + 1
        ws.cell(row, 1, line)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
        cell = ws.cell(row, 1)
        cell.alignment = Alignment(horizontal="center" if index < 2 else "left", wrap_text=True)
        if index < 2:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.append([None] * columns)
    return ws.max_row + 1


def add_cover_sheet(wb: Workbook, pages: list[dict]) -> None:
    ws = wb.create_sheet("Cover Page")
    ws.column_dimensions["A"].width = 22
    for column in "BCDEFG":
        ws.column_dimensions[column].width = 22
    for index, line in enumerate(cover_lines(pages), 1):
        ws.cell(index, 1, line)
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=7)
        cell = ws.cell(index, 1)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.font = Font(bold=index <= 6, size=12 if index <= 2 else 11, color="FFFFFF" if index <= 2 else "000000")
        if index <= 2:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.row_dimensions[index].height = 24


def build_excel(pages: list[dict], questions, destination: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    add_cover_sheet(wb, pages)
    grouped = defaultdict(list)
    for question in questions:
        grouped[question.section].append(question)
    names: set[str] = set()
    intros = section_intros(pages)
    for section, items in grouped.items():
        languages = detected_languages(items) or ["English"]
        if "MCQ" in section:
            display_name = f"Section A - MCQ ({len(items)} Ques)"
            headers = ["Q.No"]
            headers.extend(f"Question ({language})" for language in languages)
            for letter in "ABCD":
                headers.extend(f"Option {letter} ({language})" for language in languages)
            widths = [9] + [44] * len(languages) + [25] * (4 * len(languages))
        elif "Short" in section:
            display_name = f"Section B - Short Ans ({len(items)})"
            headers = ["Q.No"] + [f"Question ({language})" for language in languages] + ["Marks", "Answer"]
            widths = [9] + [46] * len(languages) + [10, 34]
        else:
            display_name = f"Section B - Long Ans ({len(items)})"
            headers = ["Q.No"] + [f"Question ({language})" for language in languages] + ["Marks", "Answer"]
            widths = [9] + [46] * len(languages) + [10, 34]
        ws = wb.create_sheet(safe_sheet_name(display_name, names))
        header_row = add_intro(ws, intros.get(section, []), len(headers))
        ws.append(headers)
        for question in items:
            question_values = question_language_values(question)
            if "MCQ" in section:
                row = [question.number]
                row.extend(readable_math(question_values.get(language, "")) for language in languages)
                for letter in "ABCD":
                    option_values = option_language_values(question, letter)
                    row.extend(readable_math(option_values.get(language, "")) for language in languages)
                ws.append(row)
            else:
                marks = 5 if "Long" in section else 2
                ws.append(
                    [question.number]
                    + [readable_math(question_values.get(language, "")) for language in languages]
                    + [marks, ""]
                )
        style_sheet(ws, header_row, widths)
    if len(wb.sheetnames) == 1:
        ws = wb.create_sheet("Questions")
        headers = ["Q.No", "Question (English)", "Option A (English)", "Option B (English)", "Option C (English)", "Option D (English)"]
        ws.append(headers)
        style_sheet(ws, 1, [9, 44, 25, 25, 25, 25])
    wb.save(destination)


def write_validation(questions, stem: str) -> dict:
    sections = defaultdict(list)
    for question in questions:
        sections[question.section].append(question)
    report = {"document": stem, "total_questions": len(questions), "sections": {}}
    for name, items in sections.items():
        numbers = [item.number for item in items]
        missing_numbers = []
        if numbers:
            present = set(numbers)
            missing_numbers = [number for number in range(min(numbers), max(numbers) + 1) if number not in present]
        report["sections"][name] = {
            "questions": len(items),
            "number_range": [min(numbers), max(numbers)] if numbers else None,
            "missing_numbers": missing_numbers,
            "missing_english": [item.number for item in items if not item.english],
            "missing_hindi": [item.number for item in items if not item.hindi],
            "incomplete_options": [
                item.number
                for item in items
                if "MCQ" in name and any(not item.option(letter) for letter in "ABCD")
            ],
        }
    EVALUATION.mkdir(parents=True, exist_ok=True)
    (EVALUATION / f"{stem}_structure.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def build_word(pages: list[dict], title: str, destination: Path) -> int:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    renderer = WordMathRenderer()
    failures = 0
    last_section = ""
    for page_index, page in enumerate(pages):
        duplicate_pool: set[str] = set()
        blocks = sorted(page.get("blocks", []), key=lambda b: b.get("reading_order", 0))
        for block in blocks:
            if block.get("skipped") or block.get("label") in {"PageHeader", "PageFooter", "Picture"}:
                continue
            text = text_of(block.get("html", ""))
            if not text:
                continue
            detected_section = section_kind(text)
            if detected_section and detected_section != last_section:
                doc.add_heading(text, level=1)
                last_section = detected_section
                continue
            if block.get("label") == "SectionHeader":
                doc.add_heading(readable_math(text), level=2)
                continue
            logical_lines = html_logical_lines(block.get("html", ""), block.get("label", ""))
            normalized = [
                re.sub(r"\s+", " ", MATH_MARKER.sub("<math>", line.text)).strip().casefold()
                for line in logical_lines
            ]
            emitted = 0
            for line, key in zip(logical_lines, normalized):
                if len(logical_lines) == 1 and key and key in duplicate_pool:
                    duplicate_pool.remove(key)
                    continue
                if re.match(r"^\d{1,3}\.\s", line.text) and key not in duplicate_pool:
                    duplicate_pool.clear()
                paragraph = doc.add_paragraph()
                failures += renderer.add_logical_line(paragraph, line)
                paragraph.paragraph_format.space_after = Pt(2)
                if re.match(r"^\([A-D]\)\s", line.text):
                    paragraph.paragraph_format.left_indent = Inches(0.25)
                elif re.match(r"^(?:i|ii|iii|iv|v|vi|vii|viii|ix|x)\)\s", line.text, re.I):
                    paragraph.paragraph_format.left_indent = Inches(0.3)
                elif re.match(r"^\d{1,3}\.\s", line.text):
                    paragraph.paragraph_format.space_before = Pt(5)
                elif line.text.strip() in {"2", "5"}:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                emitted += 1
            if len(logical_lines) > 1 and emitted:
                duplicate_pool = {key for key in normalized if key}
        if page_index < len(pages) - 1:
            doc.add_page_break()
    doc.save(destination)
    return failures


def main() -> int:
    parser = argparse.ArgumentParser(description="Create structured Word and Excel documents from saved Surya JSON")
    parser.add_argument("json", nargs="?", type=Path, help="Specific *_surya.json file; default exports all")
    parser.add_argument(
        "--destination",
        type=Path,
        default=DESTINATION,
        help="Output directory (default: output/documents)",
    )
    args = parser.parse_args()
    files = [args.json] if args.json else sorted(SOURCE.glob("*_surya.json"))
    files = [path.resolve() for path in files if path and path.exists()]
    if not files:
        print("No Surya JSON output found. Run run_surya.py first.")
        return 1
    destination = args.destination.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    for index, source in enumerate(files, 1):
        print(f"[{index}/{len(files)}] Reading preserved OCR: {source.name}")
        pages = json.loads(source.read_text(encoding="utf-8"))
        stem = source.stem.removesuffix("_surya")
        questions = parse_questions(pages)
        print(f"[{index}/{len(files)}] Parsed {len(questions)} question records; creating Word...")
        failures = build_word(pages, stem, destination / f"{stem}.docx")
        print(f"[{index}/{len(files)}] Creating section-wise Excel...")
        build_excel(pages, questions, destination / f"{stem}.xlsx")
        report = write_validation(questions, stem)
        review_count = sum(
            len(details["missing_english"])
            + len(details["missing_hindi"])
            + len(details["incomplete_options"])
            for details in report["sections"].values()
        )
        print(
            f"[{index}/{len(files)}] Completed {stem}: {len(questions)} questions, "
            f"{failures} equation(s) kept as original text, {review_count} structural review item(s)."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
