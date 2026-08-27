from __future__ import annotations

import re

import mathml2omml
from docx.text.paragraph import Paragraph
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree

MATH_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/math"
OPTION_PREFIX = re.compile(r"^\s*(\([A-Da-d]\))\s*(.*)$", re.S)
MATH_ATOM = re.compile(
    r"("
    r"\\frac\{[^{}]+\}\{[^{}]+\}|"
    r"\\sqrt(?:\[[^\]]+\])?\{[^{}]+\}|"
    r"\\overline\{[^{}]+\}|"
    r"\\(?:neq|leq|geq|times|div|pm|infty|alpha|beta|gamma|theta|pi)\b|"
    r"(?:\([^()]+\)|[A-Za-z0-9])(?:\^\{[-+A-Za-z0-9]+\}|\^[-+A-Za-z0-9])|"
    r"[A-Za-z]+_\{?[-+A-Za-z0-9]+\}?"
    r")"
)


class WordMathRenderer:
    @property
    def native_equations_available(self) -> bool:
        return True

    def _omml(self, latex: str):
        mathml = latex_to_mathml(latex)
        omml = mathml2omml.convert(mathml)
        omml = omml.replace("<m:oMath>", f'<m:oMath xmlns:m="{MATH_NAMESPACE}">', 1)
        return etree.fromstring(omml.encode("utf-8"))

    def add_equation(self, paragraph: Paragraph, latex: str) -> bool:
        try:
            paragraph._p.append(self._omml(latex.strip()))
            return True
        except Exception:
            paragraph.add_run(latex)
            return False

    def add_mixed(self, paragraph: Paragraph, text: str) -> int:
        failures = 0
        cursor = 0
        for match in MATH_ATOM.finditer(text or ""):
            if match.start() > cursor:
                paragraph.add_run(text[cursor : match.start()])
            if not self.add_equation(paragraph, match.group(0)):
                failures += 1
            cursor = match.end()
        if cursor < len(text or ""):
            paragraph.add_run((text or "")[cursor:])
        return failures

    def add_block(self, paragraph: Paragraph, text: str, equation_block: bool) -> int:
        if not equation_block:
            return self.add_mixed(paragraph, text)
        option = OPTION_PREFIX.match(text or "")
        expression = text
        if option:
            paragraph.add_run(option.group(1) + " ").bold = True
            expression = option.group(2)
        return 0 if self.add_equation(paragraph, expression) else 1
